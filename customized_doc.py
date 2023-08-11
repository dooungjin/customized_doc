from docx import Document
import subprocess

def replace_placeholders(doc, placeholders):
    for p in doc.paragraphs:
        for placeholder, replacement in placeholders.items():
            p.text = p.text.replace(placeholder, replacement)

def convert_word_to_pdf(word_file_path, pdf_file_path):
    cmd = [
        "unoconv",  # Command to run unoconv
        "-f", "pdf",  # Convert to PDF format
        "-o", pdf_file_path,  # Output PDF file path
        word_file_path,  # Input Word file path
    ]

    subprocess.run(cmd)

def main():
    template_file = 'path_to_template'  # Replace with the actual path to your template
    output_file = 'output_path.docx'  # Replace with the desired output path
    first_customization = input("Enter the first customization: ")
    second_customization = input("Enter the second customization: ")

    placeholders = {
        '[FIRST CUSTOMIZATION]': first_customization,
        '[SECOND CUSTOMIZATION]': sescond_customization
    }

    doc = Document(template_file)
    replace_placeholders(doc, placeholders)
    doc.save(output_file)
    convert_word_to_pdf('output_path.docx', 'output_path.pdf')

if __name__ == "__main__":
    main()

